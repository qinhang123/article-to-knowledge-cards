#!/usr/bin/env python3
"""
Content extraction script for the make-knowledge-cards skill.

Extracts text from TXT, Markdown, PDF, PPTX, and DOCX files.

Usage:
    python extract_content.py <input_file> [--output <output_path>]

If --output is omitted, extracted text is printed to stdout.
"""

import argparse
import sys
from pathlib import Path


def extract_txt(file_path):
    """Extract text from a plain text or markdown file."""
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def extract_pdf(file_path):
    """Extract text from a PDF file using pypdf."""
    try:
        from pypdf import PdfReader
    except ImportError:
        print(
            "ERROR: pypdf is not installed. Install it with: pip install pypdf",
            file=sys.stderr,
        )
        sys.exit(1)

    reader = PdfReader(str(file_path))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"--- Page {i + 1} ---\n{text}")
    return "\n\n".join(pages) if pages else ""


def extract_pptx(file_path):
    """Extract text from a PPTX file using python-pptx."""
    try:
        from pptx import Presentation
    except ImportError:
        print(
            "ERROR: python-pptx is not installed. Install it with: pip install python-pptx",
            file=sys.stderr,
        )
        sys.exit(1)

    prs = Presentation(str(file_path))
    slides = []
    for i, slide in enumerate(prs.slides):
        parts = [f"--- Slide {i + 1} ---"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        parts.append(text)
            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells
                    )
                    if row_text.strip():
                        parts.append(row_text)
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                parts.append(f"[Notes: {notes_text}]")
        if len(parts) > 1:
            slides.append("\n".join(parts))
    return "\n\n".join(slides) if slides else ""


def extract_docx(file_path):
    """Extract text from a DOCX file using python-docx."""
    try:
        from docx import Document
    except ImportError:
        print(
            "ERROR: python-docx is not installed. Install it with: pip install python-docx",
            file=sys.stderr,
        )
        sys.exit(1)

    doc = Document(str(file_path))
    parts = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                parts.append(row_text)
    return "\n".join(parts) if parts else ""


EXTRACTORS = {
    ".txt": extract_txt,
    ".md": extract_txt,
    ".markdown": extract_txt,
    ".pdf": extract_pdf,
    ".pptx": extract_pptx,
    ".docx": extract_docx,
}


def main():
    parser = argparse.ArgumentParser(
        description="Extract text content from various document formats."
    )
    parser.add_argument("input_file", help="Path to the input file")
    parser.add_argument(
        "--output", "-o", default=None, help="Output file path (default: stdout)"
    )
    args = parser.parse_args()

    file_path = Path(args.input_file)
    if not file_path.exists():
        print(f"ERROR: File not found: {file_path}", file=sys.stderr)
        sys.exit(1)

    ext = file_path.suffix.lower()
    extractor = EXTRACTORS.get(ext)
    if extractor is None:
        print(
            f"ERROR: Unsupported file format '{ext}'. "
            f"Supported formats: {', '.join(sorted(EXTRACTORS.keys()))}",
            file=sys.stderr,
        )
        sys.exit(1)

    text = extractor(file_path)
    if not text.strip():
        print(
            f"WARNING: No text content was extracted from {file_path}.",
            file=sys.stderr,
        )
        if ext == ".pdf":
            print(
                "The PDF may be image-based (scanned). OCR is not supported by this skill.",
                file=sys.stderr,
            )
        sys.exit(0)

    if args.output:
        output_path = Path(args.output)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"Extracted content saved to: {output_path}", file=sys.stderr)
        print(f"Character count: {len(text)}", file=sys.stderr)
    else:
        print(text)


if __name__ == "__main__":
    main()
