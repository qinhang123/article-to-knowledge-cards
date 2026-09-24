#!/usr/bin/env python3
"""Generate test files in PDF, PPTX, and DOCX formats."""

from pathlib import Path

# === Test content: same article for all formats to test extraction ===

TITLE = "Database Indexing and Query Optimization"

SECTIONS = [
    {
        "heading": "What is a Database Index",
        "body": (
            "A database index is a data structure that improves the speed of data retrieval "
            "operations on a database table. An index is created on one or more columns of a "
            "table. Without an index, the database must scan every row in the table to find "
            "matching records (a full table scan). With an index, the database can quickly "
            "locate the relevant rows. The most common index type is the B-tree index, which "
            "maintains sorted data and allows searches, insertions, and deletions in "
            "logarithmic time."
        ),
    },
    {
        "heading": "B-Tree Index Structure",
        "body": (
            "A B-tree is a self-balancing tree data structure that maintains sorted data. "
            "Each node in a B-tree can contain multiple keys and multiple child pointers. "
            "This makes B-trees shallower than binary trees for the same number of keys, "
            "which means fewer disk reads during lookups. The database reads data in fixed-size "
            "blocks called pages, and B-trees are designed so that each node fits within a "
            "single page. The root node is always kept in memory, so a typical lookup requires "
            "only a few disk reads even for very large tables."
        ),
    },
    {
        "heading": "Composite Indexes",
        "body": (
            "A composite index is an index on multiple columns. The order of columns in a "
            "composite index matters significantly. The index is sorted by the first column, "
            "then by the second column within groups of the same first column value, and so on. "
            "For example, an index on (last_name, first_name) can efficiently serve queries "
            "that filter by last_name alone, or by both last_name and first_name. However, it "
            "cannot efficiently serve queries that filter by first_name alone. This is known "
            "as the leftmost prefix rule."
        ),
    },
    {
        "heading": "Index Selectivity",
        "body": (
            "Selectivity is a measure of how many rows an index will filter out. A highly "
            "selective index filters out most rows, leaving a small result set. For example, "
            "a unique index has perfect selectivity because each value matches at most one row. "
            "The query optimizer uses selectivity to decide whether to use an index. If an "
            "index is not selective enough (for example, an index on a gender column with only "
            "two values), the optimizer may prefer a full table scan because the overhead of "
            "reading the index and then fetching the rows is higher than just scanning the table."
        ),
    },
    {
        "heading": "Covering Indexes",
        "body": (
            "A covering index is an index that contains all the columns needed by a query. "
            "When a covering index is available, the database can satisfy the query entirely "
            "from the index without accessing the table data. This is called an index-only scan. "
            "Index-only scans are significantly faster than regular index lookups because they "
            "avoid the random I/O of fetching table rows. For example, if a query selects "
            "first_name and last_name from a users table, an index on (last_name, first_name) "
            "can serve as a covering index."
        ),
    },
    {
        "heading": "Query Execution Plans",
        "body": (
            "A query execution plan describes how the database will execute a SQL statement. "
            "It includes information about which tables are accessed, which indexes are used, "
            "the join methods employed, and the estimated cost of each operation. The query "
            "optimizer generates the plan by evaluating multiple possible execution strategies "
            "and choosing the one with the lowest estimated cost. Developers can use the "
            "EXPLAIN command to view the execution plan and identify performance issues such "
            "as missing indexes or inefficient join strategies."
        ),
    },
    {
        "heading": "Index Maintenance and Trade-offs",
        "body": (
            "Indexes improve read performance but slow down write operations. Each INSERT, "
            "UPDATE, or DELETE must update all affected indexes. Therefore, adding too many "
            "indexes can degrade write performance and increase storage usage. Indexes also "
            "need maintenance: over time, indexes can become fragmented due to page splits "
            "and deletions. Fragmented indexes have lower cache hit rates and more disk I/O. "
            "Regular index rebuilding or reorganizing can restore performance. The decision "
            "to add an index should consider the ratio of reads to writes for the affected "
            "table, the query patterns, and the available storage."
        ),
    },
    {
        "heading": "Full-Text Search Indexes",
        "body": (
            "Standard B-tree indexes are not suitable for text search because they match exact "
            "values or prefixes. Full-text search uses a specialized index structure called an "
            "inverted index. An inverted index maps each word (token) to a list of documents "
            "that contain it. The text is first processed through tokenization (splitting into "
            "words), stemming (reducing to root forms), and stop word removal (removing common "
            "words like 'the' and 'is'). This allows efficient searching for keywords within "
            "large text fields."
        ),
    },
]


def generate_pdf(output_path):
    """Generate a PDF file using reportlab."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        leftMargin=2.5 * cm,
        rightMargin=2.5 * cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontSize=18,
        spaceAfter=20,
    )
    heading_style = ParagraphStyle(
        "CustomHeading",
        parent=styles["Heading2"],
        fontSize=13,
        spaceBefore=16,
        spaceAfter=8,
    )
    body_style = ParagraphStyle(
        "CustomBody",
        parent=styles["Normal"],
        fontSize=11,
        leading=18,
        spaceAfter=12,
    )

    story = []
    story.append(Paragraph(TITLE, title_style))
    story.append(Spacer(1, 12))

    for section in SECTIONS:
        story.append(Paragraph(section["heading"], heading_style))
        story.append(Paragraph(section["body"], body_style))

    doc.build(story)
    print(f"PDF generated: {output_path}")


def generate_pptx(output_path):
    """Generate a PPTX file using python-pptx."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()

    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    title_slide = prs.slides.add_slide(title_slide_layout)
    title_slide.shapes.title.text = TITLE
    subtitle = title_slide.placeholders[1]
    subtitle.text = "A comprehensive overview"

    # Content slides
    for section in SECTIONS:
        content_slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(content_slide_layout)
        title = slide.shapes.title
        title.text = section["heading"]

        body = slide.placeholders[1]
        text_frame = body.text_frame
        text_frame.text = section["body"]
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(14)

    prs.save(str(output_path))
    print(f"PPTX generated: {output_path}")


def generate_docx(output_path):
    """Generate a DOCX file using python-docx."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    # Title
    doc.add_heading(TITLE, level=0)

    # Sections
    for section in SECTIONS:
        doc.add_heading(section["heading"], level=1)
        doc.add_paragraph(section["body"])

    doc.save(str(output_path))
    print(f"DOCX generated: {output_path}")


def main():
    base_dir = Path(__file__).parent

    generate_pdf(base_dir / "database_indexing.pdf")
    generate_pptx(base_dir / "database_indexing.pptx")
    generate_docx(base_dir / "database_indexing.docx")

    print("\nAll test files generated successfully.")


if __name__ == "__main__":
    main()
